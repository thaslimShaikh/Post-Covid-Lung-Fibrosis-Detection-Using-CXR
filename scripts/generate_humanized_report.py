import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import os
import argparse

from humanized_sections import HUMANIZED_PARAGRAPHS

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def format_table(out_table, src_table, header_bg="1F4E79"):
    out_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(out_table.rows):
        src_row = src_table.rows[i]
        for j, cell in enumerate(row.cells):
            cell.text = src_row.cells[j].text.strip()
            set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)

            if i == 0:  # Header
                set_cell_shading(cell, header_bg)
                for run in p.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.name = 'Calibri'
                    run.font.size = Pt(10)
            else:  # Data
                if i % 2 == 1:
                    set_cell_shading(cell, "F2F2F2")
                for run in p.runs:
                    run.font.name = 'Calibri'
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(50, 50, 50)

TABLE_PLACEMENTS = {
    0: 34,   # Abbreviations
    1: 75,   # Lit review comparison
    2: 116,  # Tech stack
    3: 125,  # Model comparison reported
    4: 130,  # Notebook report
    5: 155,  # Risk categories
    6: 172,  # Appendix A
    7: 174,  # Appendix B
    8: 176,  # Appendix C
}

def generate_report(src_doc_path, img_dir, output_path):
    src_doc = docx.Document(src_doc_path)
    out_doc = docx.Document()

    for section in out_doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    image_placements = {
        47: os.path.join(img_dir, 'image1.png'),
        88: os.path.join(img_dir, 'image2.png'),
        127: os.path.join(img_dir, 'image3.png'),
        133: os.path.join(img_dir, 'image4.png'),
        135: os.path.join(img_dir, 'image5.png'),
        139: os.path.join(img_dir, 'image6.png'),
        153: os.path.join(img_dir, 'image7.png'),
    }

    print("Building docx document...")

    for p_idx in range(len(src_doc.paragraphs)):
        if p_idx in image_placements:
            img_path = image_placements[p_idx]
            if os.path.exists(img_path):
                img_p = out_doc.add_paragraph()
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_p.paragraph_format.space_before = Pt(12)
                img_p.paragraph_format.space_after = Pt(4)
                run = img_p.add_run()
                run.add_picture(img_path, width=Inches(5.8))

        for t_idx, pos in TABLE_PLACEMENTS.items():
            if pos == p_idx:
                src_t = src_doc.tables[t_idx]
                new_t = out_doc.add_table(rows=len(src_t.rows), cols=len(src_t.columns))
                format_table(new_t, src_t)
                t_space = out_doc.add_paragraph()
                t_space.paragraph_format.space_after = Pt(6)

        if p_idx in HUMANIZED_PARAGRAPHS:
            p_text = HUMANIZED_PARAGRAPHS[p_idx]
        else:
            p_text = src_doc.paragraphs[p_idx].text.strip()

        if not p_text:
            continue

        p = out_doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)

        if p_text.startswith("CHAPTER") or p_text in ["CERTIFICATE", "DECLARATION", "ACKNOWLEDGEMENT", "ABSTRACT", "TABLE OF CONTENTS", "LIST OF ABBREVIATIONS", "REFERENCES", "FINAL PROJECT SUMMARY"] or p_text.startswith("APPENDIX"):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(p_text)
            run.font.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(31, 78, 121)
        elif p_text.startswith("Figure "):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(p_text)
            run.font.italic = True
            run.font.name = 'Calibri'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(80, 80, 80)
        elif p_text.startswith(("1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9.")):
            parts = p_text.split(" ", 1)
            if len(parts) > 1 and parts[0].count(".") >= 1 and parts[0].replace(".", "").isdigit():
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(p_text)
                run.font.bold = True
                run.font.name = 'Calibri'
                run.font.size = Pt(12)
                run.font.color.rgb = RGBColor(46, 117, 182)
            else:
                run = p.add_run(p_text)
                run.font.name = 'Calibri'
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(40, 40, 40)
        elif p_text.startswith("Medical_Image_Processing_Project/"):
            p.paragraph_format.left_indent = Inches(0.3)
            run = p.add_run(p_text)
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(50, 50, 50)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = p.add_run(p_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(30, 30, 30)

    out_doc.save(output_path)
    print(f"Successfully saved humanized document to {output_path}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate humanized DOCX report.")
    parser.add_argument("--src", default="/tmp/file_attachments/report/PROJECT_REPORT_rewritten.docx", help="Path to source docx")
    parser.add_argument("--img-dir", default="/tmp/docx_extracted_images/word/media", help="Path to image directory")
    parser.add_argument("--out", default="PROJECT_REPORT_humanized.docx", help="Path to output docx")
    args = parser.parse_args()

    generate_report(args.src, args.img_dir, args.out)
